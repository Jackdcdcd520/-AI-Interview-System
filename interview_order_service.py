# -*- coding: utf-8 -*-
"""
面试顺序服务：读取 E:\\InterviewSystem\\Data\\interview_order.xlsx

数据安全：**只读** Excel，绝不写回、绝不修改原始文件。
容错：文件不存在 / 列名不规范 / 单元格为空 → 都不抛异常，返回结构化提示。
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "config"))
import settings  # noqa: E402

# 列名别名（中英文都兼容，防止用户 Excel 表头写法不同）
COLUMN_ALIASES = {
    "order": ["序号", "顺序", "面试顺序", "编号", "order", "no", "index"],
    "name": ["姓名", "名字", "候选人", "学生", "name"],
    "student_no": ["学号", "学籍号", "student_no", "studentid", "student_id", "sno"],
    "major": ["专业", "所学专业", "major"],
    "grade": ["年级", "届别", "班级", "grade", "class"],
    "position": ["应聘岗位", "岗位", "职位", "意向岗位", "position", "job"],
    "phone": ["手机", "手机号", "电话", "联系方式", "phone", "mobile", "tel"],
    "email": ["邮箱", "电子邮箱", "email", "mail"],
    "resume_file": ["简历", "简历文件", "简历文件名", "resume", "resume_file"],
    "interviewer": ["面试官", "考官", "interviewer"],
}


def _norm(s) -> str:
    return str(s).strip().lower().replace(" ", "").replace("　", "") if s is not None else ""


def _build_header_map(header_row) -> dict[int, str]:
    """把表头行映射成 {列索引: 标准字段名}。"""
    mapping: dict[int, str] = {}
    for idx, raw in enumerate(header_row):
        h = _norm(raw)
        if not h:
            continue
        for field, aliases in COLUMN_ALIASES.items():
            if any(h == _norm(a) or h.startswith(_norm(a)) for a in aliases):
                mapping[idx] = field
                break
    return mapping


def load_interview_order() -> dict:
    """
    返回：
      {
        "source": 文件路径,
        "exists": bool,
        "loaded": bool,
        "sheet": 工作表名,
        "columns_mapped": {标准字段: 表头原文},
        "rows": [ {order, name, student_no, ...}, ... ],
        "warnings": [...],
        "message": str,
      }
    """
    path = settings.INTERVIEW_ORDER_XLSX
    result: dict = {
        "source": str(path),
        "exists": path.is_file(),
        "loaded": False,
        "sheet": None,
        "columns_mapped": {},
        "rows": [],
        "warnings": [],
        "message": "",
    }

    if not path.is_file():
        result["message"] = (
            f"未找到面试顺序文件 {path}。"
            f"系统将继续可用：可运行 scripts/make_test_data.py 生成测试数据，"
            f"或把你的 interview_order.xlsx 放到该目录后重新加载。"
        )
        return result

    try:
        from openpyxl import load_workbook   # 只读打开

        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        result["sheet"] = ws.title

        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            result["message"] = "面试顺序文件为空（没有表头行）。"
            wb.close()
            return result

        header = list(header)
        colmap = _build_header_map(header)
        result["columns_mapped"] = {
            field: str(header[idx]) for idx, field in sorted(colmap.items())
        }

        if "name" not in colmap.values():
            # 退而求其次：把第一个非空列当姓名
            for idx, raw in enumerate(header):
                if _norm(raw) and idx not in colmap:
                    colmap[idx] = "name"
                    result["warnings"].append(
                        f"未识别到明确的“姓名”列，已将列 “{raw}” 当作姓名列使用。"
                    )
                    break
            else:
                result["message"] = "无法识别姓名列，请检查表头。"
                wb.close()
                return result

        parsed: list[dict] = []
        seq = 0
        for raw_row in rows:
            row = list(raw_row)
            if not any(v not in (None, "") for v in row):
                continue
            item: dict = {}
            for idx, field in colmap.items():
                if idx < len(row):
                    v = row[idx]
                    item[field] = v.strip() if isinstance(v, str) else v
            if not item.get("name"):
                continue
            seq += 1
            if item.get("order") in (None, ""):
                item["order"] = seq
            else:
                try:
                    item["order"] = int(float(item["order"]))
                except (TypeError, ValueError):
                    item["order"] = seq
            for k in ("student_no", "phone"):
                if isinstance(item.get(k), float):
                    item[k] = str(int(item[k]))
            parsed.append(item)

        parsed.sort(key=lambda r: (r.get("order") or 999999))
        result["rows"] = parsed
        result["loaded"] = True
        result["message"] = f"成功读取 {len(parsed)} 条面试顺序记录。"
        wb.close()
    except Exception as exc:
        result["message"] = f"读取面试顺序文件失败：{exc}"
        result["warnings"].append(str(exc))

    return result


if __name__ == "__main__":
    import json

    print(json.dumps(load_interview_order(), ensure_ascii=False, indent=2))


# =============================================================== 导入/同步
def sync_candidates_from_order() -> dict:
    """
    把面试顺序表同步进 candidates 表（用户自行导入名单的入口）：
      - 表里有、库里没有 → 新增候选人（并自动挂上 E 盘匹配到的简历）
      - 表里有、库里也有 → 更新顺序与基本信息（Excel 非空字段才覆盖；
        不动 resume_text / resume_parsed / 已有面试记录）
      - 简历标记全量重扫：之前没简历的人，若现在 Resumes 里放了他的文件 → 自动挂上
      - Excel 里删掉的人：库里保留（绝不删候选人及其面试记录）
    数据安全：Excel 只读；candidates 表只做新增/更新，不做删除。
    """
    from backend.app import db                      # 延迟导入避免循环依赖
    from backend.app.services import resume_service

    order = load_interview_order()
    result = {
        "ok": order["loaded"],
        "source": order["source"],
        "columns_mapped": order["columns_mapped"],
        "rows": 0,
        "added": [],
        "updated": 0,
        "resume_attached": [],
        "warnings": list(order["warnings"]),
        "message": "",
    }
    if not order["loaded"]:
        result["message"] = order["message"] or "面试顺序表不可读。"
        return result

    rows = order["rows"]
    result["rows"] = len(rows)

    cands = db.list_candidates()
    by_sno = {c.get("student_no"): c for c in cands if c.get("student_no")}
    by_name = {}
    for c in cands:
        by_name.setdefault(c.get("name"), c)

    for row in rows:
        name = str(row.get("name") or "").strip()
        if not name:
            continue
        raw_sno = row.get("student_no")
        sno = str(raw_sno).strip() if raw_sno not in (None, "") else None

        existing = by_sno.get(sno) if sno else None
        if existing is None:
            existing = by_name.get(name)

        # 磁盘上有没有这个人的简历（按姓名匹配）
        p = resume_service.match_resume(name)

        if existing is None:
            cid = db.upsert_candidate({
                "name": name,
                "student_no": sno,
                "major": row.get("major"),
                "grade": row.get("grade"),
                "position": row.get("position"),
                "phone": row.get("phone"),
                "email": row.get("email"),
                "resume_file": p.name if p else None,
                "resume_path": str(p) if p else None,
                "resume_exists": bool(p),
                "resume_text": None,
                "resume_parsed": None,
                "order_index": row.get("order"),
            })
            if cid:
                result["added"].append(name)
                by_name.setdefault(name, db.get_candidate(cid))
        else:
            merged = {
                "name": existing.get("name") or name,
                "student_no": sno or existing.get("student_no"),
                "major": row.get("major") or existing.get("major"),
                "grade": row.get("grade") or existing.get("grade"),
                "position": row.get("position") or existing.get("position"),
                "phone": row.get("phone") or existing.get("phone"),
                "email": row.get("email") or existing.get("email"),
                # resume_* 一律保留旧值，防止把已解析/已关联的简历冲掉
                "resume_file": existing.get("resume_file"),
                "resume_path": existing.get("resume_path"),
                "resume_exists": existing.get("resume_exists"),
                "resume_text": existing.get("resume_text"),
                "resume_parsed": existing.get("resume_parsed"),
                "order_index": row.get("order") or existing.get("order_index"),
            }
            if p and not existing.get("resume_exists"):
                # 之前没简历、现在磁盘上有了 → 自动补挂
                merged.update(resume_file=p.name, resume_path=str(p),
                              resume_exists=1)
                result["resume_attached"].append(name)
            db.upsert_candidate(merged)
            result["updated"] += 1

    added_n = len(result["added"])
    result["message"] = (
        f"同步完成：顺序表 {result['rows']} 行，新增 {added_n} 人，"
        f"更新 {result['updated']} 人，"
        f"新挂上简历 {len(result['resume_attached'])} 人。"
    )
    return result
